from flask import Flask, request, jsonify, render_template_string
import sqlite3
import random
import requests
import os
import html
from datetime import datetime
from openpyxl import Workbook, load_workbook

app = Flask(__name__)

# ============================================================
# CONFIGURATION
# ============================================================

DATABASE = "calvary_prayer.db"
EXCEL_FILE = "calvary_prayer_requests.xlsx"

BOTPRESS_WEBHOOK_URL = os.getenv(
    "BOTPRESS_WEBHOOK_URL",
    ""
)

ZAPIER_WEBHOOK_URL = os.getenv(
    "ZAPIER_WEBHOOK_URL",
    ""
)

# Demo OTP for testing
DEMO_OTP = True


# ============================================================
# TEMPORARY SESSION STORAGE
# ============================================================

sessions = {}


# ============================================================
# DATABASE
# ============================================================

def get_db():

    conn = sqlite3.connect(DATABASE)

    conn.row_factory = sqlite3.Row

    return conn


def init_database():

    conn = get_db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS prayer_requests (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            prayer_id TEXT UNIQUE NOT NULL,

            session_id TEXT,

            name TEXT NOT NULL,

            phone TEXT,

            language TEXT DEFAULT 'English',

            category TEXT NOT NULL,

            prayer_request TEXT NOT NULL,

            ai_response TEXT,

            phone_verified INTEGER DEFAULT 0,

            status TEXT DEFAULT 'New',

            created_at TEXT NOT NULL

        )
    """)

    conn.commit()

    conn.close()


init_database()

def migrate_database():
    """Add new columns to older databases without deleting existing data."""
    conn = get_db()
    columns = [
        row["name"]
        for row in conn.execute("PRAGMA table_info(prayer_requests)").fetchall()
    ]

    if "language" not in columns:
        conn.execute(
            "ALTER TABLE prayer_requests ADD COLUMN language TEXT DEFAULT 'English'"
        )

    conn.commit()
    conn.close()


migrate_database()



# ============================================================
# EXCEL STORAGE
# ============================================================

EXCEL_HEADERS = [
    "Prayer ID",
    "Session ID",
    "Name",
    "Phone",
    "Language",
    "Category",
    "Prayer Request",
    "AI Response",
    "Phone Verified",
    "Status",
    "Created At"
]


def get_excel_workbook():
    """Create or open the Excel workbook."""
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb["Prayer Requests"] if "Prayer Requests" in wb.sheetnames else wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Prayer Requests"
        ws.append(EXCEL_HEADERS)

    # Repair an existing empty worksheet.
    if ws.max_row == 1 and all(cell.value is None for cell in ws[1]):
        ws.delete_rows(1, 1)
        ws.append(EXCEL_HEADERS)

    return wb, ws


def save_prayer_to_excel(prayer):
    """Insert a new prayer request into Excel."""
    wb, ws = get_excel_workbook()

    ws.append([
        prayer.get("prayer_id", ""),
        prayer.get("session_id", ""),
        prayer.get("name", ""),
        prayer.get("phone", ""),
        prayer.get("language", "English"),
        prayer.get("category", ""),
        prayer.get("prayer_request", ""),
        prayer.get("ai_response", ""),
        prayer.get("phone_verified", 0),
        prayer.get("status", "New"),
        prayer.get("created_at", "")
    ])

    # Basic formatting.
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 16, "B": 38, "C": 22, "D": 18, "E": 12,
        "F": 18, "G": 50, "H": 65, "I": 16, "J": 25, "K": 22
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    wb.save(EXCEL_FILE)
    wb.close()


def update_prayer_in_excel(prayer_id, phone, phone_verified, status):
    """Update the existing Excel row after OTP verification."""
    if not os.path.exists(EXCEL_FILE):
        return False

    wb, ws = get_excel_workbook()

    prayer_id_column = 1

    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=prayer_id_column).value) == str(prayer_id):
            # Phone = column D (4)
            ws.cell(row=row, column=4).value = phone

            # Phone Verified = column I (9)
            ws.cell(row=row, column=9).value = phone_verified

            # Status = column J (10)
            ws.cell(row=row, column=10).value = status

            wb.save(EXCEL_FILE)
            wb.close()
            return True

    wb.close()
    return False



# ============================================================
# AI RESPONSE
# ============================================================

def prayer_ai_response(category, message, language="English"):
    responses = {
        "English": {
            "Job":
                "Thank you for sharing this with us. 🙏 "
                "May God guide your steps, open the right doors, "
                "provide the right opportunity, and give you "
                "peace and wisdom throughout your journey.",

            "Family":
                "Thank you for trusting us with your family "
                "request. 🙏 May God bring protection, peace, "
                "strength, unity, and wisdom to your family.",

            "Health":
                "Thank you for sharing this with us. 🙏 "
                "We will stand with you in prayer. May God "
                "give strength, comfort, peace, and healing "
                "to your loved one.",

            "Financial":
                "Thank you for sharing your financial concern. 🙏 "
                "May God provide wisdom, provision, strength, "
                "and peace during this situation.",

            "Relationship":
                "Thank you for opening your heart. 🙏 "
                "May God give you wisdom, peace, grace, "
                "understanding, and direction.",

            "General":
                "Thank you for sharing your prayer request. 🙏 "
                "We will stand with you in prayer. May God's "
                "peace, strength, guidance, and comfort be with you.",

            "Other":
                "Thank you for sharing what is on your heart. 🙏 "
                "We will stand with you in prayer. May God bring "
                "peace, strength, wisdom, guidance, and comfort "
                "to your situation."
        },

        "Telugu": {
            "Job":
                "మీ ప్రార్థనా అభ్యర్థనను మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏 "
                "దేవుడు మీ అడుగులను నడిపించి, సరైన తలుపులు తెరిచి, "
                "సరైన అవకాశాన్ని అందించి, మీ ప్రయాణంలో సమాధానం మరియు జ్ఞానాన్ని అనుగ్రహించుగాక.",

            "Family":
                "మీ కుటుంబ ప్రార్థనా అభ్యర్థనను మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏 "
                "దేవుడు మీ కుటుంబానికి రక్షణ, సమాధానం, బలం, ఐక్యత మరియు జ్ఞానాన్ని అనుగ్రహించుగాక.",

            "Health":
                "మీ అభ్యర్థనను మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏 "
                "మేము మీతో కలిసి ప్రార్థిస్తాము. దేవుడు మీకు మరియు మీ ప్రియమైన వారికి "
                "బలం, ఆదరణ, సమాధానం మరియు స్వస్థతను అనుగ్రహించుగాక.",

            "Financial":
                "మీ ఆర్థిక సమస్యను మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏 "
                "ఈ పరిస్థితిలో దేవుడు మీకు జ్ఞానం, సమృద్ధి, బలం మరియు సమాధానాన్ని అనుగ్రహించుగాక.",

            "Relationship":
                "మీ హృదయంలోని విషయాన్ని మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏 "
                "దేవుడు మీకు జ్ఞానం, సమాధానం, కృప, అవగాహన మరియు సరైన దిశను అనుగ్రహించుగాక.",

            "General":
                "మీ ప్రార్థనా అభ్యర్థనను మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏 "
                "మేము మీతో కలిసి ప్రార్థిస్తాము. దేవుని సమాధానం, బలం, మార్గదర్శకత్వం మరియు ఆదరణ మీతో ఉండుగాక.",

            "Other":
                "మీ హృదయంలోని విషయాన్ని మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏 "
                "మేము మీతో కలిసి ప్రార్థిస్తాము. దేవుడు మీ పరిస్థితిలో సమాధానం, బలం, "
                "జ్ఞానం, మార్గదర్శకత్వం మరియు ఆదరణను అనుగ్రహించుగాక."
        }
    }

    language = language if language in responses else "English"
    return responses[language].get(category, responses[language]["General"])


# ============================================================
# MAIN HTML
# ============================================================

HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Calvary Temple | Prayer AI</title>

<style>
* {
    box-sizing: border-box;
    margin: 0;
    padding: 0;
}

body {
    font-family: Arial, Helvetica, "Noto Sans Telugu", sans-serif;
    min-height: 100vh;
    background: linear-gradient(135deg, #f6eee4, #fffdf9);
    color: #302820;
}

.app {
    max-width: 900px;
    margin: auto;
    min-height: 100vh;
    background: #fffdf9;
    display: flex;
    flex-direction: column;
}

.header {
    display: flex;
    align-items: center;
    gap: 15px;
    padding: 22px;
    background: white;
    border-bottom: 1px solid #eadfd2;
    box-shadow: 0 3px 15px rgba(0,0,0,.05);
}

.logo {
    width: 58px;
    height: 58px;
    border-radius: 50%;
    background: linear-gradient(135deg, #8a5b32, #603819);
    color: white;
    display: flex;
    justify-content: center;
    align-items: center;
    font-size: 30px;
}

.header h1 {
    color: #633d21;
    font-size: 25px;
}

.header p {
    color: #918072;
    margin-top: 4px;
}

.language-panel {
    padding: 18px 22px;
    background: #fff8ef;
    border-bottom: 1px solid #eadfd2;
    text-align: center;
}

.language-panel h2 {
    color: #633d21;
    margin-bottom: 12px;
    font-size: 19px;
}

.language-buttons {
    display: flex;
    justify-content: center;
    gap: 12px;
    flex-wrap: wrap;
}

.language-btn {
    min-width: 150px;
    padding: 13px 20px;
    border: 1px solid #d9c7b1;
    border-radius: 12px;
    background: white;
    color: #633d21;
    font-size: 16px;
    font-weight: bold;
    cursor: pointer;
    transition: .2s;
}

.language-btn:hover,
.language-btn.active {
    background: #754722;
    color: white;
    border-color: #754722;
    transform: translateY(-1px);
}

.container {
    padding: 30px 20px;
    flex: 1;
}

.message {
    display: flex;
    gap: 10px;
    margin-bottom: 20px;
}

.bot {
    justify-content: flex-start;
}

.user {
    justify-content: flex-end;
}

.avatar {
    width: 42px;
    height: 42px;
    min-width: 42px;
    border-radius: 50%;
    background: #eee0d0;
    display: flex;
    justify-content: center;
    align-items: center;
    font-size: 21px;
}

.bubble {
    max-width: 650px;
    padding: 16px 18px;
    background: white;
    border-radius: 5px 18px 18px 18px;
    line-height: 1.55;
    box-shadow: 0 4px 15px rgba(0,0,0,.06);
}

.user .bubble {
    background: #754722;
    color: white;
    border-radius: 18px 5px 18px 18px;
}

.panel {
    margin-top: 20px;
    padding: 24px;
    background: white;
    border-radius: 18px;
    box-shadow: 0 5px 20px rgba(0,0,0,.06);
}

.panel h2 {
    color: #623c20;
    margin-bottom: 10px;
}

.panel p {
    color: #76695d;
    line-height: 1.6;
    margin-bottom: 15px;
}

.name-input,
input,
textarea {
    width: 100%;
    padding: 15px;
    border: 1px solid #d9c7b1;
    border-radius: 12px;
    outline: none;
    font-size: 17px;
}

textarea {
    min-height: 140px;
    resize: vertical;
    font-family: Arial, Helvetica, "Noto Sans Telugu", sans-serif;
}

input:focus,
textarea:focus {
    border-color: #7b4b25;
    box-shadow: 0 0 0 3px rgba(123,75,37,.10);
}

.btn {
    width: 100%;
    padding: 15px;
    margin-top: 15px;
    border: none;
    border-radius: 12px;
    background: #754722;
    color: white;
    font-size: 16px;
    font-weight: bold;
    cursor: pointer;
}

.btn:hover {
    background: #5d351a;
}

.btn.secondary {
    background: white;
    color: #754722;
    border: 1px solid #d9c6af;
}

.categories {
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 12px;
}

.category {
    background: white;
    border: 1px solid #e1d1be;
    padding: 18px 10px;
    border-radius: 14px;
    cursor: pointer;
    font-size: 27px;
    transition: .2s;
}

.category:hover {
    transform: translateY(-3px);
    border-color: #96683e;
    box-shadow: 0 5px 15px rgba(0,0,0,.08);
}

.category span {
    display: block;
    margin-top: 7px;
    font-size: 14px;
    color: #5f5144;
}

.success {
    text-align: center;
    padding: 35px;
}

.success-icon {
    font-size: 55px;
    margin-bottom: 15px;
}

.dashboard {
    display: block;
    text-align: center;
    margin-top: 25px;
    color: #754722;
    text-decoration: none;
    font-weight: bold;
}

footer {
    text-align: center;
    padding: 20px;
    color: #998a7c;
    font-size: 13px;
    border-top: 1px solid #eee3d7;
}

.hidden {
    display: none !important;
}

@media(max-width:650px) {
    .categories {
        grid-template-columns: repeat(2,1fr);
    }

    .container {
        padding: 20px 14px;
    }

    .header {
        padding: 18px;
    }

    .header h1 {
        font-size: 21px;
    }

    .language-btn {
        width: 100%;
    }
}
</style>
</head>

<body>

<div class="app">

<header class="header">
    <div class="logo">✝</div>
    <div>
        <h1>Calvary Temple</h1>
        <p id="headerSubtitle">🙏 Prayer AI</p>
    </div>
</header>

<!-- LANGUAGE SELECTION -->
<div id="languagePanel" class="language-panel">
    <h2 id="languageTitle">🌐 Select Your Language / మీ భాషను ఎంచుకోండి</h2>

    <div class="language-buttons">
        <button class="language-btn" id="englishBtn" onclick="selectLanguage('English')">
            🇬🇧 English
        </button>

        <button class="language-btn" id="teluguBtn" onclick="selectLanguage('Telugu')">
            🇮🇳 తెలుగు
        </button>
    </div>
</div>

<div class="container">

<div id="chat" class="hidden">

    <div class="message bot">
        <div class="avatar">🙏</div>

        <div class="bubble">
            <strong id="welcomeTitle">Shalom! 🙏</strong>
            <br><br>

            <span id="welcomeText">
                Welcome to <strong>Calvary Temple Prayer AI</strong>
            </span>

            <br><br>

            <span id="welcomeHonored">
                We would be honored to pray with you.
            </span>
        </div>
    </div>

</div>

<!-- NAME -->
<div id="namePanel" class="panel hidden">

    <h2 id="nameTitle">👋 What is your name?</h2>

    <p id="nameHelp">
        Please share your name with us.
    </p>

    <input
        id="name"
        class="name-input"
        type="text"
        placeholder="Enter your name"
    >

    <button class="btn" onclick="saveName()" id="continueBtn">
        Continue 🙏
    </button>

</div>

<!-- CATEGORY -->
<div id="categoryPanel" class="panel hidden">

    <h2 id="categoryTitle">🙏 How can we pray for you?</h2>

    <p id="categoryHelp">
        Choose a prayer topic.
    </p>

    <div class="categories">

        <button class="category" onclick="selectCategory('Job')">
            💼
            <span id="catJob">Job</span>
        </button>

        <button class="category" onclick="selectCategory('Family')">
            👨‍👩‍👧
            <span id="catFamily">Family</span>
        </button>

        <button class="category" onclick="selectCategory('Health')">
            ❤️
            <span id="catHealth">Health</span>
        </button>

        <button class="category" onclick="selectCategory('Financial')">
            💰
            <span id="catFinancial">Financial</span>
        </button>

        <button class="category" onclick="selectCategory('Relationship')">
            💕
            <span id="catRelationship">Relationship</span>
        </button>

        <button class="category" onclick="selectCategory('General')">
            🙏
            <span id="catGeneral">General</span>
        </button>

        <button class="category" onclick="selectCategory('Other')">
            🌿
            <span id="catOther">Other</span>
        </button>

    </div>
</div>

<!-- PRAYER REQUEST -->
<div id="requestPanel" class="panel hidden">

    <h2 id="requestTitle">📝 Tell us your prayer request</h2>

    <p id="requestHelp">
        Please share what is on your heart.
    </p>

    <textarea
        id="prayerMessage"
        placeholder="Example: My mother is not feeling well..."
    ></textarea>

    <button class="btn" onclick="submitPrayer()" id="submitBtn">
        Submit Prayer Request 🙏
    </button>

</div>

<!-- PRAYER CALL -->
<div id="callPanel" class="panel hidden">

    <h2 id="callTitle">📞 Would you like a prayer call?</h2>

    <p id="callHelp">
        A member of the Calvary Temple prayer team can call you
        and pray with you personally.
    </p>

    <button class="btn" onclick="showPhone()" id="yesCallBtn">
        Yes, Request Prayer Call
    </button>

    <button class="btn secondary" onclick="finish()" id="noCallBtn">
        No, Thank You
    </button>

</div>

<!-- PHONE -->
<div id="phonePanel" class="panel hidden">

    <h2 id="phoneTitle">📱 Phone Number</h2>

    <p id="phoneHelp">
        Enter your phone number to continue.
    </p>

    <input
        id="phone"
        type="tel"
        placeholder="+91 XXXXX XXXXX"
    >

    <button class="btn" onclick="sendOTP()" id="sendOtpBtn">
        Send OTP
    </button>

</div>

<!-- OTP -->
<div id="otpPanel" class="panel hidden">

    <h2 id="otpTitle">🔐 Verify Your Phone</h2>

    <p id="otpHelp">
        Enter the 6-digit OTP.
    </p>

    <input
        id="otp"
        type="text"
        maxlength="6"
        placeholder="Enter OTP"
    >

    <button class="btn" onclick="verifyOTP()" id="verifyOtpBtn">
        Verify OTP
    </button>

</div>

<!-- SUCCESS -->
<div id="successPanel" class="panel success hidden">

    <div class="success-icon">🙏</div>

    <h2 id="successTitle">
        Prayer Call Requested
    </h2>

    <p id="successThanks">
        Thank you for trusting Calvary Temple.
    </p>

    <p id="successSent">
        Your request has been sent to our prayer team.
    </p>

    <p id="successBless">
        May God bless you and give you peace.
    </p>

</div>

<a href="/dashboard" class="dashboard" id="dashboardLink">
    🔐 Prayer Team Dashboard
</a>

</div>

<footer id="footerText">
    © 2026 Calvary Temple · Prayer AI
</footer>

</div>

<script>

let sessionId =
    localStorage.getItem("calvary_prayer_session");

if (!sessionId) {
    sessionId = crypto.randomUUID();

    localStorage.setItem(
        "calvary_prayer_session",
        sessionId
    );
}

let userName = "";
let selectedCategory = "";
let prayerId = "";

let selectedLanguage =
    localStorage.getItem("calvary_prayer_language") || "";

const translations = {

    English: {
        headerSubtitle: "🙏 Prayer AI",

        languageTitle: "🌐 Select Your Language",
        shalom: "Shalom! 🙏",
        welcomeText: "Welcome to <strong>Calvary Temple Prayer AI</strong>",
        welcomeHonored: "We would be honored to pray with you.",

        nameTitle: "👋 What is your name?",
        nameHelp: "Please share your name with us.",
        namePlaceholder: "Enter your name",
        continue: "Continue 🙏",

        categoryTitle: "🙏 How can we pray for you?",
        categoryHelp: "Choose a prayer topic.",

        Job: "Job",
        Family: "Family",
        Health: "Health",
        Financial: "Financial",
        Relationship: "Relationship",
        General: "General",
        Other: "Other",

        requestTitle: "📝 Tell us your prayer request",
        requestHelp: "Please share what is on your heart.",
        requestPlaceholder: "Example: My mother is not feeling well...",
        submit: "Submit Prayer Request 🙏",

        callTitle: "📞 Would you like a prayer call?",
        callHelp: "A member of the Calvary Temple prayer team can call you and pray with you personally.",
        yesCall: "Yes, Request Prayer Call",
        noCall: "No, Thank You",

        phoneTitle: "📱 Phone Number",
        phoneHelp: "Enter your phone number to continue.",
        sendOtp: "Send OTP",

        otpTitle: "🔐 Verify Your Phone",
        otpHelp: "Enter the 6-digit OTP.",
        otpPlaceholder: "Enter OTP",
        verifyOtp: "Verify OTP",

        successTitle: "Prayer Call Requested",
        successThanks: "Thank you for trusting Calvary Temple.",
        successSent: "Your request has been sent to our prayer team.",
        successBless: "May God bless you and give you peace.",

        dashboard: "🔐 Prayer Team Dashboard",
        footer: "© 2026 Calvary Temple · Prayer AI",

        enterName: "Please enter your name.",
        requestRequired: "Please enter your prayer request.",
        phoneRequired: "Please enter your phone number.",
        otpRequired: "Please enter the OTP.",
        unableOtp: "Unable to send OTP.",
        verificationFailed: "Verification failed.",

        greeting: function(name) {
            return "Shalom, " + escapeHTML(name) + "! 🙏<br><br>" +
                   "Thank you for sharing your name with us.<br><br>" +
                   "How can we pray for you today?";
        },

        categoryMessage: function(name) {
            return "Thank you, " + escapeHTML(name) +
                   ". 🙏<br><br>Please tell us your prayer request.";
        },

        received: "Thank you for trusting us with this request. 🙏",

        verified: function(name) {
            return "✅ Thank you, " + escapeHTML(name) +
                   ".<br><br>Your phone number has been verified.<br><br>" +
                   "Your prayer-call request has been sent to the Calvary Temple prayer team. 🙏";
        },

        finish: function(name) {
            return "Thank you for sharing your prayer request, " +
                   escapeHTML(name) + ". 🙏<br><br>" +
                   "May God's peace, strength and guidance be with you.";
        }
    },

    Telugu: {
        headerSubtitle: "🙏 ప్రార్థనా AI",

        languageTitle: "🌐 మీ భాషను ఎంచుకోండి",
        shalom: "షాలోమ్! 🙏",
        welcomeText: "<strong>క్యాల్వరీ టెంపుల్ ప్రార్థనా AI</strong> కి స్వాగతం",
        welcomeHonored: "మీతో కలిసి ప్రార్థించడం మాకు ఆనందంగా ఉంటుంది.",

        nameTitle: "👋 మీ పేరు ఏమిటి?",
        nameHelp: "దయచేసి మీ పేరును మాతో పంచుకోండి.",
        namePlaceholder: "మీ పేరును నమోదు చేయండి",
        continue: "కొనసాగించండి 🙏",

        categoryTitle: "🙏 మీ కోసం ఎలా ప్రార్థించాలి?",
        categoryHelp: "ప్రార్థనా అంశాన్ని ఎంచుకోండి.",

        Job: "ఉద్యోగం",
        Family: "కుటుంబం",
        Health: "ఆరోగ్యం",
        Financial: "ఆర్థికం",
        Relationship: "సంబంధాలు",
        General: "సాధారణం",
        Other: "ఇతర",

        requestTitle: "📝 మీ ప్రార్థనా అభ్యర్థనను తెలియజేయండి",
        requestHelp: "మీ హృదయంలో ఉన్న విషయాన్ని మాతో పంచుకోండి.",
        requestPlaceholder: "ఉదాహరణ: మా అమ్మగారి ఆరోగ్యం బాగోలేదు...",
        submit: "ప్రార్థనా అభ్యర్థన పంపండి 🙏",

        callTitle: "📞 ప్రార్థనా కాల్ కావాలా?",
        callHelp: "క్యాల్వరీ టెంపుల్ ప్రార్థనా బృందంలోని ఒక సభ్యుడు మీకు కాల్ చేసి వ్యక్తిగతంగా మీతో కలిసి ప్రార్థించగలరు.",
        yesCall: "అవును, ప్రార్థనా కాల్ కోరుతున్నాను",
        noCall: "వద్దు, ధన్యవాదాలు",

        phoneTitle: "📱 ఫోన్ నంబర్",
        phoneHelp: "కొనసాగించడానికి మీ ఫోన్ నంబర్ నమోదు చేయండి.",
        sendOtp: "OTP పంపండి",

        otpTitle: "🔐 మీ ఫోన్‌ను ధృవీకరించండి",
        otpHelp: "6 అంకెల OTP నమోదు చేయండి.",
        otpPlaceholder: "OTP నమోదు చేయండి",
        verifyOtp: "OTP ధృవీకరించండి",

        successTitle: "ప్రార్థనా కాల్ అభ్యర్థించబడింది",
        successThanks: "క్యాల్వరీ టెంపుల్‌పై నమ్మకం ఉంచినందుకు ధన్యవాదాలు.",
        successSent: "మీ అభ్యర్థన మా ప్రార్థనా బృందానికి పంపబడింది.",
        successBless: "దేవుడు మిమ్మల్ని ఆశీర్వదించి సమాధానాన్ని అనుగ్రహించుగాక.",

        dashboard: "🔐 ప్రార్థనా బృందం డ్యాష్‌బోర్డ్",
        footer: "© 2026 క్యాల్వరీ టెంపుల్ · ప్రార్థనా AI",

        enterName: "దయచేసి మీ పేరును నమోదు చేయండి.",
        requestRequired: "దయచేసి మీ ప్రార్థనా అభ్యర్థనను నమోదు చేయండి.",
        phoneRequired: "దయచేసి మీ ఫోన్ నంబర్‌ను నమోదు చేయండి.",
        otpRequired: "దయచేసి OTP నమోదు చేయండి.",
        unableOtp: "OTP పంపడం సాధ్యం కాలేదు.",
        verificationFailed: "ధృవీకరణ విఫలమైంది.",

        greeting: function(name) {
            return "షాలోమ్, " + escapeHTML(name) +
                   "! 🙏<br><br>" +
                   "మీ పేరును మాతో పంచుకున్నందుకు ధన్యవాదాలు.<br><br>" +
                   "ఈ రోజు మీ కోసం ఎలా ప్రార్థించాలి?";
        },

        categoryMessage: function(name) {
            return "ధన్యవాదాలు, " + escapeHTML(name) +
                   ". 🙏<br><br>మీ ప్రార్థనా అభ్యర్థనను మాతో పంచుకోండి.";
        },

        received: "ఈ అభ్యర్థనను మాతో పంచుకున్నందుకు ధన్యవాదాలు. 🙏",

        verified: function(name) {
            return "✅ ధన్యవాదాలు, " + escapeHTML(name) +
                   ".<br><br>మీ ఫోన్ నంబర్ ధృవీకరించబడింది.<br><br>" +
                   "మీ ప్రార్థనా కాల్ అభ్యర్థన క్యాల్వరీ టెంపుల్ ప్రార్థనా బృందానికి పంపబడింది. 🙏";
        },

        finish: function(name) {
            return "మీ ప్రార్థనా అభ్యర్థనను పంచుకున్నందుకు ధన్యవాదాలు, " +
                   escapeHTML(name) + ". 🙏<br><br>" +
                   "దేవుని సమాధానం, బలం మరియు మార్గదర్శకత్వం మీతో ఉండుగాక.";
        }
    }
};

function t() {
    return translations[selectedLanguage] || translations.English;
}

function selectLanguage(language) {

    selectedLanguage = language;

    localStorage.setItem(
        "calvary_prayer_language",
        selectedLanguage
    );

    document.getElementById("englishBtn")
        .classList.toggle("active", language === "English");

    document.getElementById("teluguBtn")
        .classList.toggle("active", language === "Telugu");

    applyLanguage();

    document.getElementById("languagePanel")
        .classList.add("hidden");

    document.getElementById("chat")
        .classList.remove("hidden");

    document.getElementById("namePanel")
        .classList.remove("hidden");

    scrollDown();
}

function applyLanguage() {

    if (!selectedLanguage) {
        return;
    }

    const lang = t();

    document.getElementById("headerSubtitle").textContent =
        lang.headerSubtitle;

    document.getElementById("languageTitle").textContent =
        lang.languageTitle;

    document.getElementById("welcomeTitle").textContent =
        lang.shalom;

    document.getElementById("welcomeText").innerHTML =
        lang.welcomeText;

    document.getElementById("welcomeHonored").textContent =
        lang.welcomeHonored;

    document.getElementById("nameTitle").textContent =
        lang.nameTitle;

    document.getElementById("nameHelp").textContent =
        lang.nameHelp;

    document.getElementById("name").placeholder =
        lang.namePlaceholder;

    document.getElementById("continueBtn").textContent =
        lang.continue;

    document.getElementById("categoryTitle").textContent =
        lang.categoryTitle;

    document.getElementById("categoryHelp").textContent =
        lang.categoryHelp;

    document.getElementById("catJob").textContent = lang.Job;
    document.getElementById("catFamily").textContent = lang.Family;
    document.getElementById("catHealth").textContent = lang.Health;
    document.getElementById("catFinancial").textContent = lang.Financial;
    document.getElementById("catRelationship").textContent = lang.Relationship;
    document.getElementById("catGeneral").textContent = lang.General;
    document.getElementById("catOther").textContent = lang.Other;

    document.getElementById("requestTitle").textContent =
        lang.requestTitle;

    document.getElementById("requestHelp").textContent =
        lang.requestHelp;

    document.getElementById("prayerMessage").placeholder =
        lang.requestPlaceholder;

    document.getElementById("submitBtn").textContent =
        lang.submit;

    document.getElementById("callTitle").textContent =
        lang.callTitle;

    document.getElementById("callHelp").textContent =
        lang.callHelp;

    document.getElementById("yesCallBtn").textContent =
        lang.yesCall;

    document.getElementById("noCallBtn").textContent =
        lang.noCall;

    document.getElementById("phoneTitle").textContent =
        lang.phoneTitle;

    document.getElementById("phoneHelp").textContent =
        lang.phoneHelp;

    document.getElementById("sendOtpBtn").textContent =
        lang.sendOtp;

    document.getElementById("otpTitle").textContent =
        lang.otpTitle;

    document.getElementById("otpHelp").textContent =
        lang.otpHelp;

    document.getElementById("otp").placeholder =
        lang.otpPlaceholder;

    document.getElementById("verifyOtpBtn").textContent =
        lang.verifyOtp;

    document.getElementById("successTitle").textContent =
        lang.successTitle;

    document.getElementById("successThanks").textContent =
        lang.successThanks;

    document.getElementById("successSent").textContent =
        lang.successSent;

    document.getElementById("successBless").textContent =
        lang.successBless;

    document.getElementById("dashboardLink").textContent =
        lang.dashboard;

    document.getElementById("footerText").textContent =
        lang.footer;
}

async function saveName() {

    const name =
        document.getElementById("name").value.trim();

    if (!name) {
        alert(t().enterName);
        return;
    }

    userName = name;

    addUserMessage(name);

    addBotMessage(t().greeting(name));

    try {

        await fetch("/api/user", {
            method: "POST",

            headers: {
                "Content-Type": "application/json"
            },

            body: JSON.stringify({
                session_id: sessionId,
                name: name,
                language: selectedLanguage
            })
        });

    } catch(error) {
        console.error(error);
    }

    document.getElementById("namePanel")
        .classList.add("hidden");

    document.getElementById("categoryPanel")
        .classList.remove("hidden");

    scrollDown();
}

function selectCategory(category) {

    selectedCategory = category;

    addUserMessage(
        emoji(category) + " " + t()[category]
    );

    addBotMessage(
        t().categoryMessage(userName)
    );

    document.getElementById("categoryPanel")
        .classList.add("hidden");

    document.getElementById("requestPanel")
        .classList.remove("hidden");

    scrollDown();
}

async function submitPrayer() {

    const message =
        document.getElementById("prayerMessage")
            .value.trim();

    if (!message) {
        alert(t().requestRequired);
        return;
    }

    addUserMessage(message);

    document.getElementById("requestPanel")
        .classList.add("hidden");

    addBotMessage(t().received);

    try {

        const response = await fetch("/api/prayer", {

            method: "POST",

            headers: {
                "Content-Type": "application/json"
            },

            body: JSON.stringify({
                session_id: sessionId,
                category: selectedCategory,
                message: message,
                language: selectedLanguage
            })
        });

        const data = await response.json();

        if (!data.success) {

            addBotMessage(data.message);
            return;
        }

        prayerId = data.prayer_id;

        addBotMessage(data.ai_response);

        setTimeout(function() {

            document.getElementById("callPanel")
                .classList.remove("hidden");

            scrollDown();

        }, 800);

    } catch(error) {

        console.error(error);

        addBotMessage(
            selectedLanguage === "Telugu"
                ? "మీ ప్రార్థనా అభ్యర్థన స్వీకరించబడింది. 🙏"
                : "Your prayer request has been received. 🙏"
        );
    }
}

function showPhone() {

    document.getElementById("callPanel")
        .classList.add("hidden");

    document.getElementById("phonePanel")
        .classList.remove("hidden");

    scrollDown();
}

async function sendOTP() {

    const phone =
        document.getElementById("phone")
            .value.trim();

    if (!phone) {
        alert(t().phoneRequired);
        return;
    }

    try {

        const response = await fetch(
            "/api/prayer-call",
            {
                method: "POST",

                headers: {
                    "Content-Type": "application/json"
                },

                body: JSON.stringify({
                    session_id: sessionId,
                    phone: phone,
                    prayer_id: prayerId
                })
            }
        );

        const data = await response.json();

        if (!data.success) {
            alert(data.message);
            return;
        }

        if (data.demo_otp) {

            alert(
                (selectedLanguage === "Telugu"
                    ? "డెమో OTP: "
                    : "DEMO OTP: ") +
                data.demo_otp
            );
        }

        document.getElementById("phonePanel")
            .classList.add("hidden");

        document.getElementById("otpPanel")
            .classList.remove("hidden");

        scrollDown();

    } catch(error) {

        console.error(error);
        alert(t().unableOtp);
    }
}

async function verifyOTP() {

    const otp =
        document.getElementById("otp")
            .value.trim();

    if (!otp) {
        alert(t().otpRequired);
        return;
    }

    try {

        const response = await fetch(
            "/api/verify-otp",
            {
                method: "POST",

                headers: {
                    "Content-Type": "application/json"
                },

                body: JSON.stringify({
                    session_id: sessionId,
                    otp: otp
                })
            }
        );

        const data = await response.json();

        if (!data.success) {
            alert(data.message);
            return;
        }

        document.getElementById("otpPanel")
            .classList.add("hidden");

        document.getElementById("successPanel")
            .classList.remove("hidden");

        addBotMessage(
            t().verified(userName)
        );

        scrollDown();

    } catch(error) {

        console.error(error);
        alert(t().verificationFailed);
    }
}

function finish() {

    document.getElementById("callPanel")
        .classList.add("hidden");

    addBotMessage(
        t().finish(userName)
    );

    scrollDown();
}

function addUserMessage(text) {

    const chat =
        document.getElementById("chat");

    const div =
        document.createElement("div");

    div.className = "message user";

    const bubble =
        document.createElement("div");

    bubble.className = "bubble";
    bubble.textContent = text;

    div.appendChild(bubble);
    chat.appendChild(div);

    scrollDown();
}

function addBotMessage(text) {

    const chat =
        document.getElementById("chat");

    const div =
        document.createElement("div");

    div.className = "message bot";

    div.innerHTML = `
        <div class="avatar">🙏</div>
        <div class="bubble">${text}</div>
    `;

    chat.appendChild(div);

    scrollDown();
}

function scrollDown() {

    window.scrollTo({
        top: document.body.scrollHeight,
        behavior: "smooth"
    });
}

function emoji(category) {

    const icons = {
        Job: "💼",
        Family: "👨‍👩‍👧",
        Health: "❤️",
        Financial: "💰",
        Relationship: "💕",
        General: "🙏",
        Other: "🌿"
    };

    return icons[category] || "🙏";
}

function escapeHTML(text) {

    const div =
        document.createElement("div");

    div.textContent = text;

    return div.innerHTML;
}

/* Restore saved language.
   If there is no saved language, the user must select one. */
if (selectedLanguage) {
    selectLanguage(selectedLanguage);
}

</script>

</body>
</html>
"""


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():

    return render_template_string(
        HTML
    )


# ============================================================
# SAVE USER NAME
# ============================================================

@app.route(
    "/api/user",
    methods=["POST"]
)
def save_user():

    data = request.get_json() or {}


    session_id =data.get(
            "session_id"
        )

    name =data.get(
            "name"
        )

    language = data.get("language", "English")


    if not session_id or not name:

        return jsonify({

            "success":
                False,

            "message":
                "Name is required."

        }), 400


    sessions[
        session_id
    ] = {

        "name":
            name,

        "language":
            language,

        "phone":
            "",

        "otp":
            "",

        "prayer_id":
            ""

    }


    return jsonify({

        "success":
            True,

        "message":
            "Name saved."

    })


# ============================================================
# CREATE PRAYER REQUEST
# ============================================================

@app.route(
    "/api/prayer",
    methods=["POST"]
)
def create_prayer():

    data =request.get_json() or {}


    session_id =data.get(
            "session_id"
        )

    category =data.get(
            "category"
        )

    message =data.get(
            "message"
        )


    if not session_id:

        return jsonify({

            "success":
                False,

            "message":
                "Session not found."

        }), 400


    if not category or not message:

        return jsonify({

            "success":
                False,

            "message":
                "Prayer category and request are required."

        }), 400


    user =sessions.get(
            session_id
        )


    if not user:

        return jsonify({

            "success":
                False,

            "message":
                "Please enter your name first."

        }), 400


    name =user.get(
            "name",
            ""
        )

    language = data.get(
        "language",
        user.get("language", "English")
    )

    if language not in ("English", "Telugu"):
        language = "English"


    prayer_id = (

        "CT-" +

        str(
            random.randint(
                100000,
                999999
            )
        )

    )


    ai_response =prayer_ai_response(
            category,
            message,
            language
        )


    created_at =datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )


    conn =get_db()


    conn.execute(

        """
        INSERT INTO prayer_requests
        (
            prayer_id,
            session_id,
            name,
            category,
            language,
            prayer_request,
            ai_response,
            phone_verified,
            status,
            created_at
        )

        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,

        (

            prayer_id,

            session_id,

            name,

            category,

            language,

            message,

            ai_response,

            0,

            "New",

            created_at

        )

    )


    conn.commit()

    conn.close()

    # Save the same prayer request to Excel.
    save_prayer_to_excel({
        "prayer_id": prayer_id,
        "session_id": session_id,
        "name": name,
        "phone": "",
        "language": language,
        "category": category,
        "prayer_request": message,
        "ai_response": ai_response,
        "phone_verified": 0,
        "status": "New",
        "created_at": created_at
    })


    # Save prayer ID in session

    sessions[
        session_id
    ][
        "prayer_id"
    ] = prayer_id


    # Send to Botpress

    send_to_botpress({

        "prayer_id":
            prayer_id,

        "name":
            name,

        "category":
            category,

        "language":
            language,

        "prayer_request":
            message,

        "ai_response":
            ai_response,

        "status":
            "New",

        "created_at":
            created_at

    })


    return jsonify({

        "success":
            True,

        "prayer_id":
            prayer_id,

        "ai_response":
            ai_response

    })


# ============================================================
# PRAYER CALL / OTP
# ============================================================

@app.route(
    "/api/prayer-call",
    methods=["POST"]
)
def prayer_call():

    data =request.get_json() or {}


    session_id =data.get(
            "session_id"
        )

    phone =data.get(
            "phone"
        )

    prayer_id =data.get(
            "prayer_id"
        )


    if not session_id or not phone:

        return jsonify({

            "success":
                False,

            "message":
                "Phone number is required."

        }), 400


    if not prayer_id:

        return jsonify({

            "success":
                False,

            "message":
                "Prayer request not found."

        }), 400


    otp =str(
            random.randint(
                100000,
                999999
            )
        )


    sessions[
        session_id
    ][
        "phone"
    ] = phone


    sessions[
        session_id
    ][
        "otp"
    ] = otp


    print()
    print(
        "======================================"
    )
    print(
        "CALVARY TEMPLE DEMO OTP"
    )
    print(
        "Phone:",
        phone
    )
    print(
        "OTP:",
        otp
    )
    print(
        "======================================"
    )
    print()


    return jsonify({

        "success":
            True,

        "demo_otp":
            otp
            if DEMO_OTP
            else None

    })


# ============================================================
# VERIFY OTP
# ============================================================

@app.route(
    "/api/verify-otp",
    methods=["POST"]
)
def verify_otp():

    data =request.get_json() or {}


    session_id =data.get(
            "session_id"
        )

    otp = str(
            data.get(
                "otp",
                ""
            )
        )


    user =sessions.get(
            session_id
        )


    if not user:

        return jsonify({

            "success":
                False,

            "message":
                "Session not found."

        }), 400


    if otp != str(
        user.get(
            "otp"
        )
    ):

        return jsonify({

            "success":
                False,

            "message":
                "Invalid OTP."

        }), 400


    prayer_id =user.get(
            "prayer_id"
        )


    phone =user.get(
            "phone"
        )


    # --------------------------------------------------------
    # UPDATE DATABASE
    # --------------------------------------------------------

    conn =get_db()


    conn.execute(

        """
        UPDATE prayer_requests

        SET
            phone = ?,
            phone_verified = 1,
            status = 'Prayer Call Requested'

        WHERE prayer_id = ?
        """,

        (
            phone,
            prayer_id
        )

    )


    conn.commit()


    row =conn.execute(

            """
            SELECT *
            FROM prayer_requests
            WHERE prayer_id = ?
            """,

            (
                prayer_id,
            )

        ).fetchone()


    conn.close()


    # --------------------------------------------------------
    # ZAPIER
    # --------------------------------------------------------

    if row:

        send_to_zapier(
            dict(row)
        )


    return jsonify({

        "success":
            True,

        "message":
            "Phone verified."

    })


# ============================================================
# BOTPRESS
# ============================================================

def send_to_botpress(
    prayer
):

    if not BOTPRESS_WEBHOOK_URL:

        print(
            "Botpress webhook not configured."
        )

        return


    try:

        response =requests.post(

                BOTPRESS_WEBHOOK_URL,

                json=prayer,

                timeout=10

            )


        print(
            "Botpress:",
            response.status_code
        )


    except Exception as error:

        print(
            "Botpress error:",
            error
        )


# ============================================================
# ZAPIER
# ============================================================

def send_to_zapier(
    prayer
):

    if not ZAPIER_WEBHOOK_URL:

        print(
            "Zapier webhook not configured."
        )

        print(
            "Prayer Team Request:"
        )

        print(
            prayer
        )

        return


    payload = {

        "organization":
            "Calvary Temple",

        "prayer_id":
            prayer[
                "prayer_id"
            ],

        "name":
            prayer[
                "name"
            ],

        "phone":
            prayer[
                "phone"
            ],

        "category":
            prayer[
                "category"
            ],

        "language":
            prayer.get(
                "language",
                "English"
            ),

        "prayer_request":
            prayer[
                "prayer_request"
            ],

        "phone_verified":
            prayer[
                "phone_verified"
            ],

        "status":
            prayer[
                "status"
            ],

        "created_at":
            prayer[
                "created_at"
            ]

    }


    try:

        response =requests.post(

                ZAPIER_WEBHOOK_URL,

                json=payload,

                timeout=10

            )


        print(
            "Zapier:",
            response.status_code
        )


    except Exception as error:

        print(
            "Zapier error:",
            error
        )


# ============================================================
# PRAYER TEAM DASHBOARD
# ============================================================

@app.route(
    "/dashboard"
)
def dashboard():

    conn =get_db()


    rows =conn.execute(

            """
            SELECT *
            FROM prayer_requests
            ORDER BY id DESC
            """

        ).fetchall()


    conn.close()


    table_rows = ""


    for row in rows:

        verified = (

            "✅ Verified"

            if row[
                "phone_verified"
            ]

            else
            "—"

        )


        status_class = (

            "call"

            if row[
                "status"
            ] ==
            "Prayer Call Requested"

            else
            "new"

        )


        table_rows += f"""

        <tr>

            <td>
                {html.escape(
                    row["prayer_id"]
                )}
            </td>

            <td>
                {html.escape(
                    row["name"]
                    or "Anonymous"
                )}
            </td>

            <td>
                {html.escape(
                    row["phone"]
                    or "Not provided"
                )}
            </td>

            <td>
                {html.escape(
                    row["language"] or "English"
                )}
            </td>

            <td>
                {html.escape(
                    row["category"]
                )}
            </td>

            <td>
                {html.escape(
                    row["prayer_request"]
                )}
            </td>

            <td>
                {verified}
            </td>

            <td>

                <span class="status {status_class}">

                    {html.escape(
                        row["status"]
                    )}

                </span>

            </td>

            <td>
                {html.escape(
                    row["created_at"]
                )}
            </td>

        </tr>

        """


    dashboard_html = f"""

<!DOCTYPE html>

<html>

<head>

<meta charset="UTF-8">

<meta name="viewport"
      content="width=device-width, initial-scale=1.0">

<title>
Calvary Temple | Prayer Team
</title>

<style>

body {{

    font-family:
        Arial,
        Helvetica,
        sans-serif;

    background: #f6eee4;

    margin: 0;

    padding: 25px;

    color: #332920;
}}

.container {{

    max-width: 1450px;

    margin: auto;
}}

h1 {{

    color: #613c20;

    margin-bottom: 5px;
}}

.subtitle {{

    color: #88796b;

    margin-bottom: 25px;
}}

.back {{

    display: inline-block;

    margin-bottom: 20px;

    color: #754722;

    text-decoration: none;

    font-weight: bold;
}}

.card {{

    background: white;

    padding: 20px;

    border-radius: 15px;

    box-shadow:
        0 4px 15px
        rgba(0,0,0,.06);

    overflow-x: auto;
}}

table {{

    width: 100%;

    border-collapse: collapse;

    min-width: 1200px;
}}

th {{

    background: #754722;

    color: white;

    padding: 13px;

    text-align: left;
}}

td {{

    padding: 12px;

    border-bottom:
        1px solid #eee2d5;

    vertical-align: top;
}}

tr:hover {{

    background: #fff9f3;
}}

.status {{

    display: inline-block;

    padding: 6px 10px;

    border-radius: 20px;

    font-size: 12px;

    font-weight: bold;
}}

.new {{

    background: #fff0cc;

    color: #8a6100;
}}

.call {{

    background: #dff4e4;

    color: #24723a;
}}

</style>

</head>

<body>

<div class="container">

<a
    href="/"
    class="back">

    ← Back to Prayer AI

</a>

<h1>
    🙏 Calvary Temple Prayer Team
</h1>

<div class="subtitle">
    Prayer Requests
</div>

<div class="card">

<table>

<thead>

<tr>

<th>
Prayer ID
</th>

<th>
Name
</th>

<th>
Phone
</th>

<th>
Language
</th>

<th>
Category
</th>

<th>
Prayer Request
</th>

<th>
Phone
</th>

<th>
Status
</th>

<th>
Created
</th>

</tr>

</thead>

<tbody>

{table_rows}

</tbody>

</table>

</div>

</div>

</body>

</html>

"""


    return render_template_string(
        dashboard_html
    )


# ============================================================
# API: ALL PRAYER REQUESTS
# ============================================================

@app.route(
    "/api/prayer-requests"
)
def prayer_requests_api():

    conn =get_db()


    rows =conn.execute(

            """
            SELECT *
            FROM prayer_requests
            ORDER BY id DESC
            """

        ).fetchall()


    conn.close()


    return jsonify({

        "success":
            True,

        "count":
            len(rows),

        "requests":
            [
                dict(row)
                for row in rows
            ]

    })


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route(
    "/api/health"
)
def health():

    conn =get_db()


    count =conn.execute(

            """
            SELECT COUNT(*)
            FROM prayer_requests
            """

        ).fetchone()[0]


    conn.close()


    return jsonify({

        "application":
            "Calvary Temple Prayer AI",

        "status":
            "running",

        "database":
            DATABASE,

        "excel_file":
            EXCEL_FILE,

        "prayer_requests":
            count,

        "botpress":
            bool(
                BOTPRESS_WEBHOOK_URL
            ),

        "zapier":
            bool(
                ZAPIER_WEBHOOK_URL
            ),

        "demo_otp":
            DEMO_OTP

    })


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    print()
    print(
        "=========================================="
    )
    print(
        "🙏 CALVARY TEMPLE PRAYER AI"
    )
    print(
        "=========================================="
    )
    print()
    print(
        "Prayer AI:"
    )
    print(
        "http://127.0.0.1:5000"
    )
    print()
    print(
        "Prayer Team Dashboard:"
    )
    print(
        "http://127.0.0.1:5000/dashboard"
    )
    print()
    print(
        "Database:"
    )
    print(
        "calvary_prayer.db"
    )
    print()
    print(
        "Excel:"
    )
    print(
        "calvary_prayer_requests.xlsx"
    )
    print()
    print(
        "=========================================="
    )


    app.run(

        host="127.0.0.1",

        port=5000,

        debug=True

    )


if __name__ == "__main__":
    import os
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
